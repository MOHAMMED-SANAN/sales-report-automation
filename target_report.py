import pandas as pd

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter


# =============================================================
# 1. READ AND CLEAN INPUT FILE
# =============================================================

def target_newfile(file_path):

    df = pd.read_excel(
        file_path,
        header=6
    )

    # =========================================================
    # STORE NAME CLEANING
    # =========================================================

    # IMPORTANT:
    # DO NOT MERGE SPM BRANCHES WITH NORMAL BRANCHES
    #
    # Al Quoz        -> Al Quoz
    # Al Quoz(SPM)   -> Al Quoz(SPM)
    #
    # Abu_Dhabi_SPM  -> Abu_Dhabi_SPM
    # Abu_Dhabi_UAE  -> Abu_Dhabi_UAE
    #
    # Ras Al Khor(SPM) -> Ras Al Khor(SPM)
    # Dubai            -> Dubai

    df["STORE"] = (
        df["STORE"]
        .astype("string")
        .str.strip()
    )

    # =========================================================
    # REMOVE EXPORT(SPM)
    # =========================================================

    df = df[
        df["STORE"] != "Export(SPM)"
    ]

    # =========================================================
    # CONVERT DATE COLUMN
    # =========================================================

    df["DATE"] = pd.to_datetime(
        df["DATE"],
        dayfirst=True,
        errors="coerce"
    )

    return df


# =============================================================
# 2. CREATE TARGET REPORT
# =============================================================

def target_percenntage(df):

    # =========================================================
    # 1. ACHIEVED SALES
    # =========================================================

    sales = (
        df.groupby("STORE")["SUB TOTAL"]
        .sum()
    )

    # =========================================================
    # 2. FIXED STORE-WISE TARGET
    # =========================================================

    # IMPORTANT:
    # EVERY BRANCH IS SEPARATE
    #
    # NO SPM BRANCH IS MERGED WITH ANOTHER BRANCH

    targets = {

        "Abu_Dhabi_SPM": 300000,

        "Abu_Dhabi_UAE": 500000,

        "AJMAN": 400000,

        "AL QUOZ": 300000,

        "Al Quoz(SPM)": 700000,

        "DIP": 300000,

        "DUBAI": 200000,

        "Ras Al Khor(SPM)": 700000,

        "QUSAIS": 600000,

        "RASHIDIYA": 600000,

        "SHARJAH": 800000
    }

    # =========================================================
    # 3. CREATE REPORT
    # =========================================================

    report = pd.DataFrame({

        "POSITION": range(
            1,
            len(targets) + 1
        ),

        "BRANCH": list(
            targets.keys()
        ),

        "TOTAL TARGET": list(
            targets.values()
        )
    })

    # =========================================================
    # 4. ACHIEVED SALES TILL NOW
    # =========================================================

    report["ACHIEVED SALES TILL NOW"] = (

        report["BRANCH"]
        .map(sales)
        .fillna(0)
    )

    # =========================================================
    # 5. ACHIEVED SALES % TILL NOW
    # =========================================================

    report["ACHIEVED SALES % TILL NOW"] = (

        report["ACHIEVED SALES TILL NOW"]
        / report["TOTAL TARGET"]

    ).round(4)

    # =========================================================
    # 6. REMAINING TARGET
    # =========================================================

    report["REMAINING TARGET"] = (

        report["TOTAL TARGET"]
        - report["ACHIEVED SALES TILL NOW"]

    ).clip(
        lower=0
    )

    # =========================================================
    # 7. REPORT DATE
    # =========================================================

    report_date = (

        pd.to_datetime(
            df["DATE"],
            errors="coerce"
        )
        .max()
        .normalize()
    )

    # =========================================================
    # 8. MONTH START AND END
    # =========================================================

    month_start = report_date.replace(
        day=1
    )

    month_end = (

        report_date
        + pd.offsets.MonthEnd(0)
    )

    # =========================================================
    # 9. CREATE DATE RANGES
    # =========================================================

    # Full month

    all_days = pd.date_range(

        start=month_start,

        end=month_end
    )

    # 1st of month -> REPORT DATE

    days_till_now = pd.date_range(

        start=month_start,

        end=report_date
    )

    # =========================================================
    # 10. WORKING DAYS
    # =========================================================

    branch_working_days = {}

    branch_working_days_till_now = {}

    for branch in report["BRANCH"]:

        # -----------------------------------------------------
        # SHARJAH & AJMAN
        # FRIDAY = WEEKLY HOLIDAY
        # -----------------------------------------------------

        if branch in [

            "SHARJAH",

            "AJMAN"

        ]:

            working_days = (

                all_days.dayofweek != 4

            ).sum()

            working_days_till_now = (

                days_till_now.dayofweek != 4

            ).sum()

        # -----------------------------------------------------
        # OTHER BRANCHES
        # SUNDAY = WEEKLY HOLIDAY
        # -----------------------------------------------------

        else:

            working_days = (

                all_days.dayofweek != 6

            ).sum()

            working_days_till_now = (

                days_till_now.dayofweek != 6

            ).sum()

        # -----------------------------------------------------
        # STORE RESULTS
        # -----------------------------------------------------

        branch_working_days[branch] = (
            working_days
        )

        branch_working_days_till_now[branch] = (
            working_days_till_now
        )

    # =========================================================
    # 11. ADD WORKING DAYS TO REPORT
    # =========================================================

    report["WORKING DAYS THIS MONTH"] = (

        report["BRANCH"]
        .map(branch_working_days)
    )

    report["WORKING DAYS TILL NOW"] = (

        report["BRANCH"]
        .map(
            branch_working_days_till_now
        )
    )

    # =========================================================
    # 12. REMAINING DAYS
    # =========================================================

    report["REMAINING DAYS"] = (

        report["WORKING DAYS THIS MONTH"]
        - report["WORKING DAYS TILL NOW"]
    )

    # =========================================================
    # 13. REMAINING DAILY TARGET SHOULD ACHIEVE
    # =========================================================

    report[
        "REMAINING DAILY TARGET SHOULD ACHIEVE"
    ] = (

        report["REMAINING TARGET"]

        .div(
            report["REMAINING DAYS"]
        )

        .replace(
            [
                float("inf"),
                -float("inf")
            ],
            0
        )

        .fillna(0)

        .round(2)
    )

    # =========================================================
    # 14. DAILY TARGET TO BE FOLLOW
    # =========================================================

    # Keep original value for calculations

    daily_target = (

        report["TOTAL TARGET"]
        / report["WORKING DAYS THIS MONTH"]
    )

    # Display value

    report[
        "DAILY TARGET TO BE FOLLOW"
    ] = (

        daily_target.round(2)
    )

    # =========================================================
    # 15. TOTAL TARGET HAVE TO BE HAPPENED TILL NOW
    # =========================================================

    report[
        "TOTAL TARGET HAVE TO BE HAPPENED TILL NOW"
    ] = (

        report["WORKING DAYS TILL NOW"]
        * daily_target

    ).round(2)

    # =========================================================
    # 16. REQUIRED SALES %
    # =========================================================

    report["REQUIRED SALES %"] = (

        report["REMAINING TARGET"]
        / report["TOTAL TARGET"]

    ).round(4)

    # =========================================================
    # 17. RETURN REPORT
    # =========================================================

    return report


# =============================================================
# 3. FORMAT REPORT
# =============================================================

def format_report(
    report,
    output_file,
    report_date
):

    # =========================================================
    # 1. ARRANGE COLUMNS
    # =========================================================

    column_order = [

        "POSITION",

        "BRANCH",

        "TOTAL TARGET",

        "ACHIEVED SALES TILL NOW",

        "REMAINING TARGET",

        "REMAINING DAYS",

        "REMAINING DAILY TARGET SHOULD ACHIEVE",

        "ACHIEVED SALES % TILL NOW",

        "WORKING DAYS THIS MONTH",

        "DAILY TARGET TO BE FOLLOW",

        "WORKING DAYS TILL NOW",

        "TOTAL TARGET HAVE TO BE HAPPENED TILL NOW",

        "REQUIRED SALES %"
    ]

    report.columns = (

        report.columns
        .astype(str)
        .str.strip()
    )

    report = report[
        column_order
    ].copy()

    # =========================================================
    # 2. SAVE DATA TO EXCEL
    # =========================================================

    report.to_excel(

        output_file,

        index=False,

        startrow=2
    )

    wb = load_workbook(
        output_file
    )

    ws = wb.active

    # =========================================================
    # 3. DATE HEADER
    # =========================================================

    date_text = report_date.strftime(
        "%Y %B %d"
    ).upper()

    last_col = len(
        report.columns
    )

    # =========================================================
    # DATE ROW COLOR
    # =========================================================

    date_fill = PatternFill(

        fill_type="solid",

        fgColor="F4B183"
    )

    # =========================================================
    # MERGE DATE ROW
    # =========================================================

    ws.merge_cells(

        start_row=1,

        start_column=1,

        end_row=1,

        end_column=last_col
    )

    # =========================================================
    # APPLY COLOR TO DATE ROW
    # =========================================================

    for col in range(
        1,
        last_col + 1
    ):

        cell = ws.cell(

            row=1,

            column=col
        )

        cell.fill = date_fill

    # =========================================================
    # DATE CELL
    # =========================================================

    date_cell = ws.cell(

        row=1,

        column=1
    )

    date_cell.value = date_text

    date_cell.font = Font(

        name="Arial",

        bold=True,

        size=14,

        color="FF0000"
    )

    date_cell.alignment = Alignment(

        horizontal="center",

        vertical="center"
    )

    # =========================================================
    # 4. COLORS
    # =========================================================

    yellow_fill = PatternFill(

        fill_type="solid",

        fgColor="FFFF00"
    )

    light_yellow_fill = PatternFill(

        fill_type="solid",

        fgColor="FFF2CC"
    )

    white_fill = PatternFill(

        fill_type="solid",

        fgColor="FFFFFF"
    )

    total_orange_fill = PatternFill(

        fill_type="solid",

        fgColor="F4B183"
    )

    # =========================================================
    # 5. BORDERS
    # =========================================================

    orange_side = Side(

        style="thin",

        color="F4B183"
    )

    black_side = Side(

        style="thin",

        color="000000"
    )

    # =========================================================
    # NORMAL TABLE BORDER
    # =========================================================

    normal_border = Border(

        left=orange_side,

        right=orange_side,

        top=orange_side,

        bottom=orange_side
    )

    # =========================================================
    # TOTAL ROW BORDER
    # =========================================================

    total_border = Border(

        top=Side(

            style="medium",

            color="000000"
        ),

        bottom=Side(

            style="thin",

            color="000000"
        )
    )

    # =========================================================
    # 6. HEADER
    # =========================================================

    header_row = 3

    for col in range(

        1,

        last_col + 1
    ):

        cell = ws.cell(

            row=header_row,

            column=col
        )

        cell.fill = yellow_fill

        cell.font = Font(

            name="Arial",

            bold=True,

            size=11
        )

        cell.alignment = Alignment(

            horizontal="center",

            vertical="center",

            wrap_text=True
        )

        cell.border = normal_border

    # =========================================================
    # 7. DATA ROWS
    # =========================================================

    first_data_row = 4

    last_data_row = (

        first_data_row

        + len(report)

        - 1
    )

    for row in range(

        first_data_row,

        last_data_row + 1
    ):

        # -----------------------------------------------------
        # ALTERNATE ROW COLORS
        # -----------------------------------------------------

        if (

            (row - first_data_row) % 2

            == 0

        ):

            fill = light_yellow_fill

        else:

            fill = white_fill

        # -----------------------------------------------------
        # APPLY TO ALL CELLS
        # -----------------------------------------------------

        for col in range(

            1,

            last_col + 1
        ):

            cell = ws.cell(

                row=row,

                column=col
            )

            cell.fill = fill

            cell.border = normal_border

            # ALL VALUES BOLD

            cell.font = Font(

                name="Arial",

                bold=True,

                size=11
            )

            cell.alignment = Alignment(

                horizontal="center",

                vertical="center"
            )

    # =========================================================
    # 8. NUMBER FORMATTING
    # =========================================================

    for col_num, column_name in enumerate(

        report.columns,

        start=1
    ):

        column_name = str(
            column_name
        ).upper()

        for row in range(

            first_data_row,

            last_data_row + 1
        ):

            cell = ws.cell(

                row=row,

                column=col_num
            )

            # -------------------------------------------------
            # PERCENTAGE COLUMNS
            # -------------------------------------------------

            if "%" in column_name:

                cell.number_format = "0.00%"

            # -------------------------------------------------
            # TARGET / SALES COLUMNS
            # -------------------------------------------------

            elif any(

                word in column_name

                for word in [

                    "TARGET",

                    "SALES"

                ]

            ):

                cell.number_format = "#,##0.00"

            # -------------------------------------------------
            # DAYS COLUMNS
            # -------------------------------------------------

            elif "DAYS" in column_name:

                cell.number_format = "0"

    # =========================================================
    # 9. TOTAL ROW
    # =========================================================

    total_row = (

        last_data_row

        + 1
    )

    # =========================================================
    # TOTAL LABEL
    # =========================================================

    ws.cell(

        row=total_row,

        column=2
    ).value = "TOTAL"

    # =========================================================
    # GRAND TOTAL CALCULATIONS
    # =========================================================

    total_target = (

        report[
            "TOTAL TARGET"
        ].sum()
    )

    total_achieved = (

        report[
            "ACHIEVED SALES TILL NOW"
        ].sum()
    )

    total_remaining = (

        report[
            "REMAINING TARGET"
        ].sum()
    )

    total_remaining_daily = (

        report[
            "REMAINING DAILY TARGET SHOULD ACHIEVE"
        ].sum()
    )

    total_achieved_percentage = (

        total_achieved

        / total_target
    )

    total_target_happened = (

        report[
            "TOTAL TARGET HAVE TO BE HAPPENED TILL NOW"
        ].sum()
    )

    total_required_percentage = (

        total_remaining

        / total_target
    )

    # =========================================================
    # 10. WRITE TOTAL VALUES
    # =========================================================

    # ---------------------------------------------------------
    # C - TOTAL TARGET
    # ---------------------------------------------------------

    ws.cell(

        row=total_row,

        column=3
    ).value = total_target

    ws.cell(

        row=total_row,

        column=3
    ).number_format = "#,##0.00"

    # ---------------------------------------------------------
    # D - ACHIEVED SALES
    # ---------------------------------------------------------

    ws.cell(

        row=total_row,

        column=4
    ).value = total_achieved

    ws.cell(

        row=total_row,

        column=4
    ).number_format = "#,##0.00"

    # ---------------------------------------------------------
    # E - REMAINING TARGET
    # ---------------------------------------------------------

    ws.cell(

        row=total_row,

        column=5
    ).value = total_remaining

    ws.cell(

        row=total_row,

        column=5
    ).number_format = "#,##0.00"

    # ---------------------------------------------------------
    # F - REMAINING DAYS
    # ---------------------------------------------------------

    ws.cell(

        row=total_row,

        column=6
    ).value = None

    # ---------------------------------------------------------
    # G - REMAINING DAILY TARGET
    # ---------------------------------------------------------

    ws.cell(

        row=total_row,

        column=7
    ).value = total_remaining_daily

    ws.cell(

        row=total_row,

        column=7
    ).number_format = "#,##0.00"

    # ---------------------------------------------------------
    # H - ACHIEVED SALES %
    # ---------------------------------------------------------

    ws.cell(

        row=total_row,

        column=8
    ).value = total_achieved_percentage

    ws.cell(

        row=total_row,

        column=8
    ).number_format = "0.00%"

    # ---------------------------------------------------------
    # I - WORKING DAYS THIS MONTH
    # ---------------------------------------------------------

    ws.cell(

        row=total_row,

        column=9
    ).value = None

    # ---------------------------------------------------------
    # J - DAILY TARGET
    # ---------------------------------------------------------

    ws.cell(

        row=total_row,

        column=10
    ).value = None

    # ---------------------------------------------------------
    # K - WORKING DAYS TILL NOW
    # ---------------------------------------------------------

    ws.cell(

        row=total_row,

        column=11
    ).value = None

    # ---------------------------------------------------------
    # L - TARGET HAPPENED TILL NOW
    # ---------------------------------------------------------

    ws.cell(

        row=total_row,

        column=12
    ).value = total_target_happened

    ws.cell(

        row=total_row,

        column=12
    ).number_format = "#,##0.00"

    # ---------------------------------------------------------
    # M - REQUIRED SALES %
    # ---------------------------------------------------------

    ws.cell(

        row=total_row,

        column=13
    ).value = total_required_percentage

    ws.cell(

        row=total_row,

        column=13
    ).number_format = "0.00%"

    # =========================================================
    # 11. STYLE TOTAL ROW
    # =========================================================

    for col in range(

        1,

        last_col + 1
    ):

        cell = ws.cell(

            row=total_row,

            column=col
        )

        cell.fill = light_yellow_fill

        cell.font = Font(

            name="Arial",

            bold=True,

            italic=True,

            size=11
        )

        cell.border = total_border

        cell.alignment = Alignment(

            horizontal="center",

            vertical="center"
        )

    # =========================================================
    # 12. SPECIAL TOTAL CELL
    # ACHIEVED SALES TOTAL - ORANGE
    # =========================================================

    ws.cell(

        row=total_row,

        column=4
    ).fill = total_orange_fill

    # =========================================================
    # 13. COLUMN WIDTH
    # =========================================================

    widths = {

        "POSITION": 10,

        "BRANCH": 22,

        "TOTAL TARGET": 18,

        "ACHIEVED SALES TILL NOW": 22,

        "REMAINING TARGET": 20,

        "REMAINING DAYS": 15,

        "REMAINING DAILY TARGET SHOULD ACHIEVE": 28,

        "ACHIEVED SALES % TILL NOW": 20,

        "WORKING DAYS THIS MONTH": 20,

        "DAILY TARGET TO BE FOLLOW": 25,

        "WORKING DAYS TILL NOW": 20,

        "TOTAL TARGET HAVE TO BE HAPPENED TILL NOW": 30,

        "REQUIRED SALES %": 20
    }

    for col_num, column_name in enumerate(

        report.columns,

        start=1
    ):

        column_letter = get_column_letter(
            col_num
        )

        width = widths.get(

            str(column_name).upper(),

            18
        )

        ws.column_dimensions[
            column_letter
        ].width = width

    # =========================================================
    # 14. HIDE COLUMNS
    # =========================================================

    # I = WORKING DAYS THIS MONTH
    # J = DAILY TARGET TO BE FOLLOW
    # K = WORKING DAYS TILL NOW
    # L = TOTAL TARGET HAVE TO BE HAPPENED TILL NOW

    ws.column_dimensions[
        "I"
    ].hidden = True

    ws.column_dimensions[
        "J"
    ].hidden = True

    ws.column_dimensions[
        "K"
    ].hidden = True

    ws.column_dimensions[
        "L"
    ].hidden = True

    # =========================================================
    # 15. ROW HEIGHT
    # =========================================================

    # DATE ROW

    ws.row_dimensions[
        1
    ].height = 25

    # BLANK ROW BETWEEN DATE AND HEADER

    ws.row_dimensions[
        2
    ].height = 10

    # HEADER

    ws.row_dimensions[
        3
    ].height = 60

    # DATA

    for row in range(

        first_data_row,

        last_data_row + 1
    ):

        ws.row_dimensions[
            row
        ].height = 30

    # TOTAL

    ws.row_dimensions[
        total_row
    ].height = 30

    # =========================================================
    # 16. ALIGNMENT
    # =========================================================

    # ---------------------------------------------------------
    # DATE
    # ---------------------------------------------------------

    ws["A1"].alignment = Alignment(

        horizontal="center",

        vertical="center"
    )

    # ---------------------------------------------------------
    # HEADER
    # ---------------------------------------------------------

    for cell in ws[3]:

        cell.alignment = Alignment(

            horizontal="center",

            vertical="center",

            wrap_text=True
        )

    # ---------------------------------------------------------
    # DATA + TOTAL
    # ---------------------------------------------------------

    for row in range(

        first_data_row,

        total_row + 1
    ):

        for col in range(

            1,

            last_col + 1
        ):

            ws.cell(

                row=row,

                column=col

            ).alignment = Alignment(

                horizontal="center",

                vertical="center"
            )

    # =========================================================
    # 17. FREEZE HEADER
    # =========================================================

    ws.freeze_panes = "A4"

    # =========================================================
    # 18. SAVE
    # =========================================================

    wb.save(
        output_file
    )

    print(
        "Report created successfully:"
    )

    print(
        output_file
    )


# =============================================================
# 19. RUN REPORT
# =============================================================

if __name__ == "__main__":

    # ---------------------------------------------------------
    # INPUT / OUTPUT FILE
    # ---------------------------------------------------------

    input_file = (
        "/content/sample_data/input.xlsx"
    )

    output_file = (
        "/content/target_report.xlsx"
    )

    # ---------------------------------------------------------
    # READ AND CLEAN INPUT EXCEL
    # ---------------------------------------------------------

    df = target_newfile(
        input_file
    )

    # ---------------------------------------------------------
    # CREATE TARGET REPORT
    # ---------------------------------------------------------

    report = target_percenntage(
        df
    )

    # ---------------------------------------------------------
    # USE LATEST DATE AVAILABLE
    # ---------------------------------------------------------

    report_date = (

        pd.to_datetime(
            df["DATE"],
            errors="coerce"
        )
        .max()
        .normalize()
    )

    # ---------------------------------------------------------
    # CREATE FORMATTED EXCEL REPORT
    # ---------------------------------------------------------

    format_report(

        report,

        output_file,

        report_date
    )
