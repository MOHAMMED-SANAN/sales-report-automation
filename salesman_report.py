import pandas as pd


from openpyxl import Workbook


from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment,
    Border,
    Side
)


def salesman_newfile(file_path):

    df = pd.read_excel(
        file_path,
        header=6
    )

    # ==========================================================
    # CLEAN COLUMN NAMES
    # ==========================================================

    df.columns = (
        df.columns
        .astype(str)
        .str.strip()
    )

    # ==========================================================
    # STORE MAPPING
    # ==========================================================

    store_mapping = {

        "Al Quoz(SPM)": "Al Quoz",

        "Abu_Dhabi_SPM": "Abu_Dhabi_UAE",

        "Ras Al Khor(SPM)": "Dubai"

    }

    df["STORE"] = df["STORE"].replace(
        store_mapping
    )

    # ==========================================================
    # REMOVE EXPORT(SPM)
    # ==========================================================

    df = df[
        df["STORE"] != "Export(SPM)"
    ]

    # ==========================================================
    # CLEAN STORE
    # ==========================================================

    df["STORE"] = (
        df["STORE"]
        .astype(str)
        .str.strip()
    )

    # ==========================================================
    # CLEAN SALES PERSON
    # ==========================================================

    df["SALES MAN"] = (
        df["SALESPERSON"]
        .astype(str)
        .str.strip()
        .str.upper()
    )

    # ==========================================================
    # DATE
    # ==========================================================

    df["DATE"] = pd.to_datetime(
        df["DATE"],
        dayfirst=True,
        errors="coerce"
    )

    # ==========================================================
    # SUB TOTAL
    # ==========================================================

    df["SUB TOTAL"] = pd.to_numeric(
        df["SUB TOTAL"],
        errors="coerce"
    ).fillna(0)

    return df


def sales_person_report(df):

    # ==========================================================
    # SALES PERSON TARGETS
    # ==========================================================

    sales_person_targets = {

        # ======================================================
        # ABU DHABI
        # ======================================================

        ("Abu_Dhabi_UAE", "DILEEP KUMAR"): 50000,

        ("Abu_Dhabi_UAE", "MOHAMMED NISAR"): 200000,

        ("Abu_Dhabi_UAE", "MOHAMED NISAR"): 200000,

        ("Abu_Dhabi_UAE", "MUHAMMED"): 200000,

        ("Abu_Dhabi_UAE", "MUHAMED"): 200000,

        ("Abu_Dhabi_UAE", "RENNY"): 50000,

        ("Abu_Dhabi_UAE", "SUMESH MOHAN"): 45000,

        ("Abu_Dhabi_UAE", "PRAJITH"): 80000,


        # ======================================================
        # AL QUOZ
        # ======================================================

        ("Al Quoz", "ANOOB"): 55000,

        ("Al Quoz", "DENNY"): 50000,

        ("Al Quoz", "JOYSTON LOBO"): 250000,

        ("Al Quoz", "LATHESH"): 225000,

        ("Al Quoz", "MIDHUN"): 80000,

        ("Al Quoz", "NAVEEN"): 50000,

        ("Al Quoz", "RAJESH R G"): 80000,

        ("Al Quoz", "RAMAN"): 150000,


        # ======================================================
        # DIP
        # ======================================================

        ("DIP", "LINSON WILSON"): 75000,

        ("DIP", "RAKESH V RAJAN"): 50000,


        # ======================================================
        # DUBAI
        # ======================================================

        ("Dubai", "BASIL ZAMAN"): 80000,

        ("Dubai", "ISHAQ"): 50000,

        ("Dubai", "MUJEEBU RAHMAN"): 50000,

        ("Dubai", "RAHUL RAVINDRAN"): 50000,

        ("Dubai", "SATISH"): 50000,

        ("Dubai", "SATHISH"): 50000,

        ("Dubai", "SHAN FAYAZ"): 50000,


        # ======================================================
        # QUSAIS
        # ======================================================

        ("Qusais_Sales SPM", "AKHIL"): 50000,

        ("Qusais_Sales SPM", "ARUN KUMAR H"): 65000,

        ("Qusais_Sales SPM", "JITHIN GEORGE"): 50000,

        ("Qusais_Sales SPM", "SYAMLAL"): 75000,


        # ======================================================
        # RASHIDIYA
        # ======================================================

        ("Rashidiya(SPM)", "ANEESH P"): 80000,

        ("Rashidiya(SPM)", "ANTONY CHAKO"): 50000,

        ("Rashidiya(SPM)", "ANTONY CHACKO"): 50000,

        ("Rashidiya(SPM)", "JOMON VINCENT"): 65000,

        ("Rashidiya(SPM)", "NISHAD JOSE"): 60000,

        ("Rashidiya(SPM)", "RAMAN"): None,

        ("Rashidiya(SPM)", "SIJO SKARIA"): 270000,


        # ======================================================
        # SHARJAH
        # ======================================================

        ("Sharjah_UAE", "NAVEED MUHAMMAD"): 260000,

        ("Sharjah_UAE", "NIBIN"): 275000,

        ("Sharjah_UAE", "RAMAN"): 150000,

        ("Sharjah_UAE", "SAJEER"): 200000,


        # ======================================================
        # AJMAN
        # ======================================================

        ("Ajman(SPM)", "LOK RAJ"): 200000,

        ("Ajman(SPM)", "RAMAN"): None

    }


    # ==========================================================
    # BRANCH ORDER
    # ==========================================================

    branch_order = [

        "Abu_Dhabi_UAE",

        "Al Quoz",

        "DIP",

        "Dubai",

        "Qusais_Sales SPM",

        "Rashidiya(SPM)",

        "Sharjah_UAE",

        "Ajman(SPM)"

    ]


    # ==========================================================
    # REPORT DATE
    # ==========================================================

    report_date = (
        pd.to_datetime(
            df["DATE"],
            errors="coerce"
        )
        .max()
        .normalize()
    )


    # ==========================================================
    # MONTH
    # ==========================================================

    month_start = report_date.replace(
        day=1
    )

    month_end = (
        report_date
        + pd.offsets.MonthEnd(0)
    )


    # ==========================================================
    # DATE RANGE
    # ==========================================================

    all_days = pd.date_range(
        start=month_start,
        end=month_end
    )

    days_till_now = pd.date_range(
        start=month_start,
        end=report_date
    )


    # ==========================================================
    # WORKING DAYS
    # ==========================================================

    branch_working_days = {}

    branch_working_days_till_now = {}


    for branch in branch_order:

        # ------------------------------------------------------
        # FRIDAY HOLIDAY
        # SHARJAH / AJMAN
        # ------------------------------------------------------

        if branch in [
            "Sharjah_UAE",
            "Ajman(SPM)"
        ]:

            working_days = (
                all_days.dayofweek != 4
            ).sum()

            working_days_till_now = (
                days_till_now.dayofweek != 4
            ).sum()

        # ------------------------------------------------------
        # SUNDAY HOLIDAY
        # OTHER BRANCHES
        # ------------------------------------------------------

        else:

            working_days = (
                all_days.dayofweek != 6
            ).sum()

            working_days_till_now = (
                days_till_now.dayofweek != 6
            ).sum()


        branch_working_days[
            branch
        ] = working_days


        branch_working_days_till_now[
            branch
        ] = working_days_till_now


    # ==========================================================
    # SALES
    # ==========================================================

    sales = (

        df.groupby(
            [
                "STORE",
                "SALES MAN"
            ],
            as_index=False
        )["SUB TOTAL"]

        .sum()

        .rename(
            columns={
                "SUB TOTAL":
                "ACHIEVED SALES"
            }
        )
    )


    # ==========================================================
    # REPORT ROWS
    # ==========================================================

    report_rows = []


    for branch in branch_order:

        branch_sales = sales[
            sales["STORE"] == branch
        ].copy()


        if branch_sales.empty:

            continue


        branch_sales = (
            branch_sales
            .sort_values(
                "SALES MAN"
            )
        )


        position = 1


        for _, row in branch_sales.iterrows():

            sales_person = row[
                "SALES MAN"
            ]

            achieved_sales = row[
                "ACHIEVED SALES"
            ]


            # ==================================================
            # TARGET
            # ==================================================

            monthly_target = (
                sales_person_targets.get(
                    (
                        branch,
                        sales_person
                    ),
                    None
                )
            )


            # ==================================================
            # WORKING DAYS
            # ==================================================

            working_days = (
                branch_working_days[
                    branch
                ]
            )


            working_days_till_now = (
                branch_working_days_till_now[
                    branch
                ]
            )


            # ==================================================
            # REMAINING DAYS
            # ==================================================

            remaining_days = (
                working_days
                - working_days_till_now
            )


            # ==================================================
            # CALCULATIONS
            # ==================================================

            if monthly_target is not None:

                remaining_target = (
                    monthly_target
                    - achieved_sales
                )


                if monthly_target != 0:

                    sales_percentage = (
                        achieved_sales
                        / monthly_target
                    )

                    required_sales_percentage = (
                        remaining_target
                        / monthly_target
                    )

                else:

                    sales_percentage = 0

                    required_sales_percentage = 0


                if working_days != 0:

                    daily_target = (
                        monthly_target
                        / working_days
                    )

                else:

                    daily_target = 0


                if remaining_days != 0:

                    remaining_daily_target = (
                        remaining_target
                        / remaining_days
                    )

                else:

                    remaining_daily_target = 0


                total_target_till_now = (
                    working_days_till_now
                    * daily_target
                )


            else:

                remaining_target = None

                sales_percentage = None

                required_sales_percentage = None

                daily_target = None

                remaining_daily_target = None

                total_target_till_now = None


            # ==================================================
            # ADD ROW
            # ==================================================

            report_rows.append({

                "BRANCH":
                branch,

                "SL.NO":
                position,

                "SALES MAN":
                sales_person,

                "MONTHLY TARGET":
                monthly_target,

                "ACHIEVED SALES":
                achieved_sales,

                "REMAINING TARGET":
                remaining_target,

                "REMAINING DAYS":
                remaining_days,

                "REMAINING DAILY TARGET":
                remaining_daily_target,

                "SALES %":
                sales_percentage,

                "WORKING DAYS":
                working_days,

                "DAILY TARGET":
                daily_target,

                "WORKING DAYS TILL NOW":
                working_days_till_now,

                "TOTAL TARGET":
                total_target_till_now,

                "REQUIRED SALES %":
                required_sales_percentage

            })


            position += 1


    report = pd.DataFrame(
        report_rows
    )


    return report, report_date


def set_column_widths(ws):

    widths = {

        "A": 8,

        "B": 28,

        "C": 20,

        "D": 20,

        "E": 22,

        "F": 18,

        "G": 28,

        "H": 14,

        "I": 17,

        "J": 16,

        "K": 23,

        "L": 18,

        "M": 18

    }


    for column, width in widths.items():

        ws.column_dimensions[
            column
        ].width = width


    # ==========================================================
    # HIDE I-L
    # ==========================================================

    ws.column_dimensions["I"].hidden = True

    ws.column_dimensions["J"].hidden = True

    ws.column_dimensions["K"].hidden = True

    ws.column_dimensions["L"].hidden = True


def write_headers(ws):

    headers = [

        "SL.NO",

        "SALES MAN",

        "MONTHLY TARGET",

        "ACHIEVED SALES",

        "REMAINING TARGET",

        "REMAINING DAYS",

        "REMAINING DAILY TARGET",

        "SALES %",

        "WORKING DAYS",

        "DAILY TARGET",

        "WORKING DAYS TILL NOW",

        "TOTAL TARGET",

        "REQUIRED SALES %"

    ]


    for col_num, header in enumerate(
        headers,
        start=1
    ):

        cell = ws.cell(
            row=2,
            column=col_num
        )


        cell.value = header


        cell.font = Font(
            name="Calibri",
            size=10,
            bold=True
        )


        cell.fill = PatternFill(
            "solid",
            fgColor="F4CCCC"
        )


        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )


    ws.row_dimensions[
        2
    ].height = 30


thin_border = Border(

    left=Side(
        style="thin",
        color="000000"
    ),

    right=Side(
        style="thin",
        color="000000"
    ),

    top=Side(
        style="thin",
        color="000000"
    ),

    bottom=Side(
        style="thin",
        color="000000"
    )

)


def write_sales_row(
    ws,
    row_number,
    row
):

    values = [

        row["SL.NO"],

        row["SALES MAN"],

        row["MONTHLY TARGET"],

        row["ACHIEVED SALES"],

        row["REMAINING TARGET"],

        row["REMAINING DAYS"],

        row["REMAINING DAILY TARGET"],

        row["SALES %"],

        row["WORKING DAYS"],

        row["DAILY TARGET"],

        row["WORKING DAYS TILL NOW"],

        row["TOTAL TARGET"],

        row["REQUIRED SALES %"]

    ]


    for col_num, value in enumerate(
        values,
        start=1
    ):

        cell = ws.cell(
            row=row_number,
            column=col_num
        )


        cell.value = value


        cell.border = thin_border


        # ======================================================
        # BOLD
        # ======================================================

        cell.font = Font(
            name="Calibri",
            size=10,
            bold=True
        )


        # ======================================================
        # ALIGNMENT
        # ======================================================

        if col_num == 2:

            cell.alignment = Alignment(
                horizontal="left",
                vertical="center"
            )

        else:

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )


        # ======================================================
        # MONEY FORMAT
        # ======================================================

        if col_num in [
            3,
            4,
            5,
            7,
            10,
            12
        ]:

            cell.number_format = (
                '#,##0.00'
            )


        # ======================================================
        # PERCENTAGE FORMAT
        # ======================================================

        if col_num in [
            8,
            13
        ]:

            cell.number_format = (
                '0.00%'
            )


def write_total_row(
    ws,
    row_number,
    branch_report
):

    # ==========================================================
    # TOTAL LABEL
    # ==========================================================

    cell = ws.cell(
        row=row_number,
        column=2
    )

    cell.value = "TOTAL"

    cell.font = Font(
        name="Calibri",
        size=10,
        bold=True,
        italic=True,
        color="C65911"
    )

    cell.alignment = Alignment(
        horizontal="left",
        vertical="center"
    )


    # ==========================================================
    # TOTAL TARGET
    # ==========================================================

    total_target = (
        branch_report[
            "MONTHLY TARGET"
        ]
        .dropna()
        .sum()
    )


    target_cell = ws.cell(
        row=row_number,
        column=3
    )

    target_cell.value = total_target

    target_cell.number_format = (
        '#,##0.00'
    )

    target_cell.font = Font(
        name="Calibri",
        size=10,
        bold=True,
        color="C65911"
    )


    # ==========================================================
    # TOTAL ACHIEVED
    # ==========================================================

    total_achieved = (
        branch_report[
            "ACHIEVED SALES"
        ]
        .sum()
    )


    achieved_cell = ws.cell(
        row=row_number,
        column=4
    )

    achieved_cell.value = total_achieved

    achieved_cell.number_format = (
        '#,##0.00'
    )

    achieved_cell.font = Font(
        name="Calibri",
        size=10,
        bold=True,
        color="4472C4"
    )


    # ==========================================================
    # FORMAT TOTAL ROW
    # ==========================================================

    for col in range(
        1,
        14
    ):

        cell = ws.cell(
            row=row_number,
            column=col
        )

        cell.border = thin_border

        cell.fill = PatternFill(
            "solid",
            fgColor="FFFF00"
        )


        if col not in [
            2,
            3,
            4
        ]:

            cell.font = Font(
                name="Calibri",
                size=10,
                bold=True
            )


def create_new_target_sheet(
    ws,
    report,
    report_date
):

    # ==========================================================
    # TITLE
    # ==========================================================

    title = (
        "SALES MAN DAILY TARGET REVIEW - TILL "
        + report_date.strftime(
            "%Y %B %d"
        )
    )


    ws.merge_cells(
        "A1:M1"
    )


    ws["A1"] = title


    ws["A1"].font = Font(
        name="Calibri",
        size=16,
        bold=True
    )


    ws["A1"].alignment = Alignment(
        horizontal="center",
        vertical="center"
    )


    ws["A1"].fill = PatternFill(
        "solid",
        fgColor="FFFF00"
    )


    ws.row_dimensions[
        1
    ].height = 30


    # ==========================================================
    # HEADERS
    # ==========================================================

    write_headers(ws)


    current_row = 3


    # ==========================================================
    # BRANCH DISPLAY NAMES
    # ==========================================================

    branch_display_names = {

        "Abu_Dhabi_UAE":
        "ABU DHABI",

        "Al Quoz":
        "AL QUOZ",

        "DIP":
        "DIP",

        "Dubai":
        "DUBAI",

        "Qusais_Sales SPM":
        "QUSAIS",

        "Rashidiya(SPM)":
        "RASHIDIYA",

        "Sharjah_UAE":
        "SHARJAH",

        "Ajman(SPM)":
        "AJMAN"

    }


    # ==========================================================
    # BRANCH ORDER
    # ==========================================================

    branch_order = [

        "Abu_Dhabi_UAE",

        "Al Quoz",

        "DIP",

        "Dubai",

        "Qusais_Sales SPM",

        "Rashidiya(SPM)",

        "Sharjah_UAE",

        "Ajman(SPM)"

    ]


    # ==========================================================
    # WRITE EACH BRANCH
    # ==============================================================

    for branch in branch_order:

        branch_report = report[
            report["BRANCH"] == branch
        ].copy()


        if branch_report.empty:

            continue


        # ======================================================
        # GREEN BRANCH HEADER
        # ======================================================

        ws.merge_cells(

            start_row=current_row,

            start_column=1,

            end_row=current_row,

            end_column=13

        )


        branch_cell = ws.cell(
            row=current_row,
            column=1
        )


        branch_cell.value = (
            branch_display_names[
                branch
            ]
        )


        branch_cell.font = Font(
            name="Calibri",
            size=11,
            bold=True,
            color="C65911"
        )


        branch_cell.fill = PatternFill(
            "solid",
            fgColor="C6E0B4"
        )


        branch_cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )


        current_row += 1


        # ======================================================
        # SALES PERSONS
        # ======================================================

        for _, row in branch_report.iterrows():

            write_sales_row(

                ws,

                current_row,

                row

            )

            current_row += 1


        # ======================================================
        # TOTAL
        # ======================================================

        write_total_row(

            ws,

            current_row,

            branch_report

        )


        current_row += 1


    # ==========================================================
    # WIDTHS
    # ==========================================================

    set_column_widths(ws)


    # ==========================================================
    # ROW HEIGHTS
    # ==========================================================

    for row in range(
        3,
        ws.max_row + 1
    ):

        ws.row_dimensions[
            row
        ].height = 20


    # ==========================================================
    # FREEZE
    # ==========================================================

    ws.freeze_panes = "A3"


def create_branch_sheet(
    ws,
    branch_report,
    report_date,
    branch_display_name
):

    # ==========================================================
    # TITLE
    # ==========================================================

    title = (
        branch_display_name
        + "  "
        + report_date.strftime(
            "%Y %B"
        )
    )


    ws.merge_cells(
        "A1:M1"
    )


    ws["A1"] = title


    ws["A1"].font = Font(
        name="Calibri",
        size=16,
        bold=True
    )


    ws["A1"].alignment = Alignment(
        horizontal="center",
        vertical="center"
    )


    ws["A1"].fill = PatternFill(
        "solid",
        fgColor="FFFF00"
    )


    ws.row_dimensions[
        1
    ].height = 30


    # ==========================================================
    # HEADERS
    # ==========================================================

    write_headers(ws)


    current_row = 3


    # ==========================================================
    # SALES PERSONS
    # ==========================================================

    for _, row in branch_report.iterrows():

        write_sales_row(

            ws,

            current_row,

            row

        )

        current_row += 1


    # ==========================================================
    # TOTAL
    # ==========================================================

    write_total_row(

        ws,

        current_row,

        branch_report

    )


    # ==========================================================
    # WIDTHS
    # ==========================================================

    set_column_widths(ws)


    # ==========================================================
    # ROW HEIGHT
    # ==========================================================

    for row in range(
        3,
        ws.max_row + 1
    ):

        ws.row_dimensions[
            row
        ].height = 20


    # ==========================================================
    # FREEZE
    # ==========================================================

    ws.freeze_panes = "A3"


def create_excel_report(
    report,
    report_date,
    output_file="Sales_Man_Daily_Target_Review.xlsx"
):

    # ==========================================================
    # CREATE WORKBOOK
    # ==============================================================

    wb = Workbook()


    # ==========================================================
    # ==========================================================
    # FIRST SHEET = NEW TARGET
    # ==========================================================
    # ==========================================================
    #
    # THIS SHEET IS THE SAME STYLE AS YOUR SCREENSHOT.
    #
    # ALL BRANCHES ARE TOGETHER.
    # ==============================================================

    ws = wb.active

    ws.title = "New Target"


    create_new_target_sheet(

        ws,

        report,

        report_date

    )


    # ==========================================================
    # BRANCH SHEET NAMES
    # ==========================================================

    branch_display_names = {

        "Abu_Dhabi_UAE":
        "ABU DHABI",

        "Al Quoz":
        "AL QUOZ",

        "DIP":
        "DIP",

        "Dubai":
        "DUBAI",

        "Qusais_Sales SPM":
        "QUSAIS",

        "Rashidiya(SPM)":
        "RASHIDIYA",

        "Sharjah_UAE":
        "SHARJAH",

        "Ajman(SPM)":
        "AJMAN"

    }


    # ==========================================================
    # CREATE INDIVIDUAL BRANCH SHEETS
    # ==========================================================

    for branch, sheet_name in branch_display_names.items():

        branch_report = report[
            report["BRANCH"] == branch
        ].copy()


        if branch_report.empty:

            continue


        branch_ws = wb.create_sheet(
            title=sheet_name
        )


        create_branch_sheet(

            branch_ws,

            branch_report,

            report_date,

            sheet_name

        )


    # ==========================================================
    # FORCE SHEET ORDER
    # ==========================================================

    desired_order = [

        "New Target",

        "ABU DHABI",

        "AL QUOZ",

        "DIP",

        "DUBAI",

        "QUSAIS",

        "RASHIDIYA",

        "SHARJAH",

        "AJMAN"

    ]


    wb._sheets = [

        wb[name]

        for name in desired_order

        if name in wb.sheetnames

    ]


    # ==========================================================
    # SAVE
    # ==========================================================

    wb.save(
        output_file
    )


    print()
    print(
        "=============================================="
    )

    print(
        "REPORT CREATED SUCCESSFULLY"
    )

    print(
        "=============================================="
    )

    print(
        "File:",
        output_file
    )

    print()

    print(
        "Sheets:"
    )

    for sheet in wb.sheetnames:

        print(
            "  ",
            sheet
        )
