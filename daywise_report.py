import pandas as pd


from openpyxl import load_workbook


from openpyxl.styles import Font, PatternFill, Border, Side, Alignment


from openpyxl.utils import get_column_letter


def split_date_into_columns(file_path, output_file):

    # ==========================================================
    # 1. READ EXCEL
    # ==========================================================

    df = pd.read_excel(
        file_path,
        header=6
    )

    # Clean column names
    df.columns = (
        df.columns
        .astype(str)
        .str.strip()
    )

    print("Columns found:", df.columns.tolist())

    # ==========================================================
    # 2. STORE MAPPING
    # ==========================================================

    store_mapping = {
        'Al Quoz(SPM)': 'Al Quoz',
        'Abu_Dhabi_SPM': 'Abu_Dhabi_UAE',
        'Ras Al Khor(SPM)': 'Dubai'
    }

    # Remove Export
    df = df[
        df["STORE"] != "Export(SPM)"
    ].copy()

    # Store mapping
    df["STORE"] = df["STORE"].replace(
        store_mapping
    )

    # ==========================================================
    # 3. CONVERT DATE
    # ==========================================================

    df["DATE"] = pd.to_datetime(
        df["DATE"],
        dayfirst=True,
        errors="coerce"
    )

    # Remove invalid dates
    df = df.dropna(
        subset=["DATE"]
    )

    # ==========================================================
    # 4. SUB TOTAL AS NUMBER
    # ==========================================================

    df["SUB TOTAL"] = pd.to_numeric(
        df["SUB TOTAL"],
        errors="coerce"
    ).fillna(0)

    # ==========================================================
    # 5. CREATE DATE-WISE REPORT
    # ==========================================================

    date_report = df.pivot_table(
        index="STORE",
        columns="DATE",
        values="SUB TOTAL",
        aggfunc="sum",
        fill_value=0
    )

    # ==========================================================
    # 6. SORT DATES
    # ==========================================================

    date_report = date_report.sort_index(
        axis=1
    )

    # ==========================================================
    # 7. RESET INDEX
    # ==========================================================

    date_report = date_report.reset_index()

    # Rename STORE → BRANCH
    date_report.rename(
        columns={
            "STORE": "BRANCH"
        },
        inplace=True
    )

    # ==========================================================
    # 8. ADD POSITION
    # ==========================================================

    date_report.insert(
        0,
        "POSITION",
        range(
            1,
            len(date_report) + 1
        )
    )

    # ==========================================================
    # 9. IDENTIFY DATE COLUMNS
    # ==========================================================

    date_columns = [
        col
        for col in date_report.columns
        if col not in [
            "POSITION",
            "BRANCH"
        ]
    ]

    # ==========================================================
    # 10. DAY WISE TOTAL
    # ==========================================================

    date_report["DAY WISE TOTAL"] = (
        date_report[
            date_columns
        ].sum(axis=1)
    )

    # ==========================================================
    # 11. GRAND TOTAL ROW
    # ==========================================================

    total_row = {}

    total_row["POSITION"] = ""
    total_row["BRANCH"] = "TOTAL"

    # Total for each date
    for col in date_columns:

        total_row[col] = (
            date_report[col].sum()
        )

    # Overall total
    total_row["DAY WISE TOTAL"] = (
        date_report[
            "DAY WISE TOTAL"
        ].sum()
    )

    # Add total row
    date_report.loc[
        len(date_report)
    ] = total_row

    # ==========================================================
    # 12. FORMAT DATE COLUMN NAMES
    # ==========================================================

    new_columns = []

    for col in date_report.columns:

        if hasattr(
            col,
            "strftime"
        ):

            new_columns.append(
                col.strftime("%b-%d")
            )

        else:

            new_columns.append(col)

    date_report.columns = new_columns

    # ==========================================================
    # 13. SAVE INITIAL EXCEL
    # ==========================================================

    date_report.to_excel(
        output_file,
        index=False,
        startrow=2
    )

    # ==========================================================
    # 14. OPEN EXCEL
    # ==========================================================

    wb = load_workbook(
        output_file
    )

    ws = wb.active

    # ==========================================================
    # 15. BASIC INFORMATION
    # ==========================================================

    last_column = ws.max_column

    header_row = 3

    first_data_row = 4

    last_data_row = (
        ws.max_row - 1
    )

    total_row_number = ws.max_row

    # ==========================================================
    # 16. TITLE
    # ==========================================================

    ws.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=last_column
    )

    ws.cell(
        row=1,
        column=1
    ).value = "DAYWISE SALES REPORT"

    ws.cell(
        row=1,
        column=1
    ).font = Font(
        name="Arial",
        size=12,
        bold=True,
        italic=True
    )

    ws.cell(
        row=1,
        column=1
    ).alignment = Alignment(
        horizontal="left",
        vertical="center"
    )

    # ==========================================================
    # 17. COLORS
    # ==========================================================

    yellow_fill = PatternFill(
        fill_type="solid",
        fgColor="FFD966"
    )

    light_yellow_fill = PatternFill(
        fill_type="solid",
        fgColor="FFF2CC"
    )

    green_fill = PatternFill(
        fill_type="solid",
        fgColor="A9D18E"
    )

    total_fill = PatternFill(
        fill_type="solid",
        fgColor="FCE4D6"
    )

    white_fill = PatternFill(
        fill_type="solid",
        fgColor="FFFFFF"
    )

    # ==========================================================
    # 18. FONTS
    # ==========================================================

    header_font = Font(
        name="Arial",
        size=11,
        bold=True,
        italic=True
    )

    normal_font = Font(
        name="Arial",
        size=10,
        bold=True
    )

    red_font = Font(
        name="Arial",
        size=10,
        bold=True,
        color="FF0000"
    )

    # ==========================================================
    # 19. BORDERS
    # ==========================================================

    thin_side = Side(
        style="thin",
        color="000000"
    )

    border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side
    )

    # ==========================================================
    # 20. HEADER ROW
    # ==========================================================

    for cell in ws[header_row]:

        cell.fill = yellow_fill

        cell.font = header_font

        cell.border = border

        cell.alignment = Alignment(
            horizontal="left",
            vertical="center"
        )

    # ==========================================================
    # 21. DATA ROWS
    # ==========================================================

    for row in range(
        first_data_row,
        last_data_row + 1
    ):

        # Alternate row color
        if (
            row - first_data_row
        ) % 2 == 0:

            fill = light_yellow_fill

        else:

            fill = white_fill

        for col in range(
            1,
            last_column + 1
        ):

            cell = ws.cell(
                row=row,
                column=col
            )

            cell.fill = fill

            cell.border = border

            cell.font = normal_font

            # Branch left aligned
            if col == 2:

                cell.alignment = Alignment(
                    horizontal="left",
                    vertical="center"
                )

            else:

                cell.alignment = Alignment(
                    horizontal="right",
                    vertical="center"
                )

    # ==========================================================
    # 22. NUMBER FORMATTING
    # ==========================================================

    for row in range(
        first_data_row,
        total_row_number + 1
    ):

        for col in range(
            3,
            last_column + 1
        ):

            cell = ws.cell(
                row=row,
                column=col
            )

            # Zero shown as "-"
            cell.number_format = (
                '#,##0.00;'
                '[Red]-#,##0.00;'
                '-'
            )

    # ==========================================================
    # 23. DAY WISE TOTAL COLUMN
    # ==========================================================

    total_column = last_column

    for row in range(
        header_row,
        total_row_number + 1
    ):

        cell = ws.cell(
            row=row,
            column=total_column
        )

        cell.border = border

        if row == header_row:

            cell.fill = yellow_fill

            cell.font = header_font

        else:

            cell.font = red_font

            if (
                row - first_data_row
            ) % 2 == 0:

                cell.fill = light_yellow_fill

            else:

                cell.fill = white_fill

    # ==========================================================
    # 24. TOTAL ROW
    # ==========================================================

    for col in range(
        1,
        last_column + 1
    ):

        cell = ws.cell(
            row=total_row_number,
            column=col
        )

        # Total row background
        cell.fill = total_fill

        # Border
        cell.border = border

        # ALL TOTAL ROW VALUES RED
        cell.font = Font(
            name="Arial",
            size=10,
            bold=True,
            color="FF0000"
        )

        # Alignment
        if col == 2:

            cell.alignment = Alignment(
                horizontal="left",
                vertical="center"
            )

        else:

            cell.alignment = Alignment(
                horizontal="right",
                vertical="center"
            )

    # ==========================================================
    # 25. EMPTY POSITION IN TOTAL ROW
    # ==========================================================

    ws.cell(
        row=total_row_number,
        column=1
    ).value = ""

    # ==========================================================
    # 26. GREEN ROW BELOW TOTAL
    # ==========================================================

    green_row = (
        total_row_number + 1
    )

    for col in range(
        1,
        last_column
    ):

        cell = ws.cell(
            row=green_row,
            column=col
        )

        cell.fill = green_fill

        cell.border = border

    # ==========================================================
    # 27. COLUMN WIDTHS
    # ==========================================================

    # POSITION
    ws.column_dimensions["A"].width = 12

    # BRANCH
    ws.column_dimensions["B"].width = 18

    # Date columns
    for col in range(
        3,
        last_column
    ):

        column_letter = (
            get_column_letter(col)
        )

        ws.column_dimensions[
            column_letter
        ].width = 16

    # DAY WISE TOTAL - WIDER
    ws.column_dimensions[
        get_column_letter(last_column)
    ].width = 22

    # ==========================================================
    # 28. ROW HEIGHTS
    # ==========================================================

    ws.row_dimensions[1].height = 25

    ws.row_dimensions[
        header_row
    ].height = 28

    for row in range(
        first_data_row,
        total_row_number + 1
    ):

        ws.row_dimensions[
            row
        ].height = 24

    # ==========================================================
    # 29. FREEZE PANES
    # ==========================================================

    ws.freeze_panes = "C4"

    # ==========================================================
    # 30. SAVE FINAL EXCEL
    # ==========================================================

    wb.save(
        output_file
    )

    print(
        "Excel report created successfully!"
    )

    print(
        "File:",
        output_file
    )

    return date_report
